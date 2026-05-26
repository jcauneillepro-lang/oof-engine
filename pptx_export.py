"""
OOF Engine — PPTX export
v1: render HTML to PDF (weasyprint) → split to per-page PNG (pdf2image)
    → embed each as a full-bleed image in a python-pptx 16:9 deck.

Visually perfect, not editable. Editable PPTX (v2) requires structured JSON
from the LLM and native python-pptx assembly — roadmap.
"""
import io
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
import weasyprint


SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
DPI = 144  # render quality for the PNG slides (144 = 2x at 72)


def html_to_pptx(html: str) -> bytes:
    """
    Convert a v23 HTML brief into a 16:9 PPTX.
    Returns the PPTX file as bytes (caller streams it as a download).
    """
    # Step 1 — render the HTML to a PDF (one page per slide, at 1280×720)
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = Path(tmpdir) / "brief.pdf"
        # weasyprint respects @page CSS in the HTML; our v23 print styles set
        # @page { size: 1280px 720px; margin: 0; } and page-break-after on .slide
        weasyprint.HTML(string=html).write_pdf(pdf_path)

        # Step 2 — convert each PDF page to a PNG
        # Use pdf2image (needs poppler installed in the env)
        from pdf2image import convert_from_path
        images = convert_from_path(str(pdf_path), dpi=DPI)

        # Step 3 — build the PPTX
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W_IN)
        prs.slide_height = Inches(SLIDE_H_IN)
        blank_layout = prs.slide_layouts[6]

        for img in images:
            img_path = Path(tmpdir) / f"slide_{id(img)}.png"
            img.save(img_path, "PNG")
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(img_path), 0, 0,
                width=Inches(SLIDE_W_IN), height=Inches(SLIDE_H_IN)
            )

        # Step 4 — write to bytes
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf.read()


def html_to_pdf(html: str) -> bytes:
    """Render the HTML brief to PDF (server-side, no browser needed)."""
    buf = io.BytesIO()
    weasyprint.HTML(string=html).write_pdf(buf)
    buf.seek(0)
    return buf.read()
